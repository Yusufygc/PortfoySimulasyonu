# src/application/services/excel_formatter.py

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.application.services.reporting import excel_theme as theme


class ExcelFormatter:
    def apply_formatting(self, writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
        """Profesyonel Excel formatlaması, durumsuz (stateless) operasyonlar."""
        if df.empty:
            return

        worksheet = writer.sheets[sheet_name]
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        self._apply_header_formatting(worksheet, df, thin_border)
        self._apply_zebra_rows_formatting(worksheet, df, thin_border)
        self._apply_number_and_date_formatting(worksheet, df)
        self._apply_conditional_coloring(worksheet, df)
        self._apply_summary_rows_highlighting(worksheet, df)
        self._apply_column_widths(worksheet, df)

        # BAŞLIK SATIRI DONDUR
        worksheet.freeze_panes = "A2"

        # OTOMATİK FİLTRE
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}1"

    def _apply_header_formatting(self, worksheet, df: pd.DataFrame, thin_border: Border) -> None:
        header_fill = PatternFill(start_color=theme.NAVY_PRIMARY, end_color=theme.NAVY_PRIMARY, fill_type="solid")
        header_font = Font(bold=True, color=theme.WHITE, size=11)
        for col_num, _ in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

    def _apply_zebra_rows_formatting(self, worksheet, df: pd.DataFrame, thin_border: Border) -> None:
        light_gray = PatternFill(start_color=theme.ZEBRA, end_color=theme.ZEBRA, fill_type="solid")
        for row_num in range(2, len(df) + 2):
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                if row_num % 2 == 0:
                    cell.fill = light_gray
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")

    def _apply_number_and_date_formatting(self, worksheet, df: pd.DataFrame) -> None:
        for col_num, col_name in enumerate(df.columns, 1):
            for row_num in range(2, len(df) + 2):
                cell = worksheet.cell(row=row_num, column=col_num)
                if cell.value is None:
                    continue
                if "Tarih" in col_name:
                    cell.number_format = 'dd.mm.yyyy'
                elif "(TL)" in col_name:
                    cell.number_format = '#,##0.00'
                elif "(%)" in col_name:
                    cell.number_format = '0.00%'
                elif col_name in ["Adet", "Son Adet", "Lot", "Toplam Gün Sayısı", "Aktif Pozisyon Sayısı"]:
                    cell.number_format = '#,##0'

    def _apply_conditional_coloring(self, worksheet, df: pd.DataFrame) -> None:
        green_fill = PatternFill(start_color=theme.POSITIVE_FILL, end_color=theme.POSITIVE_FILL, fill_type="solid")
        green_font = Font(color=theme.POSITIVE_FONT, bold=True)
        red_fill = PatternFill(start_color=theme.NEGATIVE_FILL, end_color=theme.NEGATIVE_FILL, fill_type="solid")
        red_font = Font(color=theme.NEGATIVE_FONT, bold=True)

        for col_num, col_name in enumerate(df.columns, 1):
            if "K/Z" in col_name or "Getiri" in col_name:
                for row_num in range(2, len(df) + 2):
                    cell = worksheet.cell(row=row_num, column=col_num)
                    if cell.value is not None and isinstance(cell.value, (int, float)):
                        if cell.value > 0:
                            cell.fill = green_fill
                            cell.font = green_font
                        elif cell.value < 0:
                            cell.fill = red_fill
                            cell.font = red_font

    def _apply_summary_rows_highlighting(self, worksheet, df: pd.DataFrame) -> None:
        bold_font = Font(bold=True, size=11)
        summary_fill = PatternFill(start_color=theme.SUMMARY_FILL, end_color=theme.SUMMARY_FILL, fill_type="solid")

        for row_num in range(2, len(df) + 2):
            ticker_cell = None
            for col_num, col_name in enumerate(df.columns, 1):
                if "Hisse" in col_name or "Ticker" in col_name:
                    ticker_cell = worksheet.cell(row=row_num, column=col_num)
                    break

            if ticker_cell and ticker_cell.value and ("TOPLAM" in str(ticker_cell.value) or "▼" in str(ticker_cell.value)):
                for col_num in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=row_num, column=col_num)
                    cell.font = bold_font
                    cell.fill = summary_fill

    def _apply_column_widths(self, worksheet, df: pd.DataFrame) -> None:
        for idx, col in enumerate(df.columns, 1):
            max_length = len(str(col))
            for row_num in range(2, min(len(df) + 2, 100)):
                cell_value = worksheet.cell(row=row_num, column=idx).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))

            adjusted_width = min(max_length + 3, 50)
            worksheet.column_dimensions[get_column_letter(idx)].width = adjusted_width
