# src/application/services/reporting/excel_layout_manager.py

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter

from src.application.services.reporting import excel_theme as theme
from src.application.services.reporting.daily_history_models import SheetName
from src.application.services.reporting.excel_data_preparer import ExcelDataPreparer


def _is_kpi_pl_label(label_cell) -> bool:
    label = str(label_cell.value or "")
    return "K/Z" in label or "Getiri" in label


def _apply_kpi_value_style(value_cell, label_cell) -> None:
    label_val = str(value_cell.value) if value_cell.value is not None else ""
    is_negative = "-" in label_val or "−" in label_val
    is_positive = not is_negative and any(c.isdigit() for c in label_val)
    if _is_kpi_pl_label(label_cell):
        if is_negative:
            value_cell.fill = PatternFill(start_color=theme.NEGATIVE_FILL, end_color=theme.NEGATIVE_FILL, fill_type="solid")
            value_cell.font = Font(bold=True, color=theme.NEGATIVE_FONT, size=13)
        elif is_positive:
            value_cell.fill = PatternFill(start_color=theme.POSITIVE_FILL, end_color=theme.POSITIVE_FILL, fill_type="solid")
            value_cell.font = Font(bold=True, color=theme.POSITIVE_FONT, size=13)
        else:
            value_cell.font = Font(bold=True, size=13)
    else:
        value_cell.font = Font(bold=True, color=theme.NAVY_DEEP, size=13)


class ExcelLayoutManager:
    def __init__(self, data_preparer: ExcelDataPreparer) -> None:
        self.data_preparer = data_preparer

    _TAB_COLORS = {
        SheetName.DASHBOARD:     "FF0D2B6E",
        SheetName.SUMMARY:       "FF00897B",
        SheetName.DAILY_DETAIL:  "FF455A64",
        SheetName.STOCK_SUMMARY: "FF1565C0",
    }

    def post_process_sheets(self, workbook) -> None:
        for sheet_name, argb in self._TAB_COLORS.items():
            if sheet_name in workbook.sheetnames:
                ws = workbook[sheet_name]
                ws.sheet_properties.tabColor = Color(rgb=argb)
                ws.page_setup.orientation = "landscape"
                ws.page_setup.fitToPage = True
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0
                ws.oddFooter.center.text = (
                    "&İ Gizli — Yalnızca İç Kullanım &İ"
                    "        &Sayfa &S / &N"
                )

    def style_dashboard_kpi(self, worksheet, dashboard_df: pd.DataFrame) -> None:
        if dashboard_df.empty:
            return
        kpi_rows = len(dashboard_df)
        for row_idx in range(2, kpi_rows + 2):
            worksheet.row_dimensions[row_idx].height = 28
            label_cell = worksheet.cell(row=row_idx, column=1)
            value_cell = worksheet.cell(row=row_idx, column=2)
            label_cell.font = Font(bold=True, color=theme.NAVY_PRIMARY, size=10)
            _apply_kpi_value_style(value_cell, label_cell)

    def add_banner_to_data_sheet(self, worksheet, title: str, ncols: int) -> None:
        if ncols <= 0:
            return
        worksheet.insert_rows(1)
        last_col = get_column_letter(max(ncols, 1))
        worksheet.merge_cells(f"A1:{last_col}1")
        banner = worksheet["A1"]
        banner.value = title
        banner.font = Font(bold=True, color=theme.WHITE, size=13)
        banner.fill = PatternFill(start_color=theme.NAVY_PRIMARY, end_color=theme.NAVY_PRIMARY, fill_type="solid")
        banner.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 26
        worksheet.freeze_panes = "A3"
        worksheet.auto_filter.ref = f"A2:{last_col}2"

