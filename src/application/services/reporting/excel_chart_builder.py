# src/application/services/reporting/excel_chart_builder.py

import pandas as pd
from openpyxl.chart import Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.application.services.reporting import excel_theme as theme
from src.application.services.reporting.daily_history_models import SheetName
from src.application.services.reporting.excel_chart_factory import ExcelChartFactory
from src.application.services.reporting.excel_data_preparer import ExcelDataPreparer


class ExcelChartBuilder:
    def __init__(self, data_preparer: ExcelDataPreparer) -> None:
        self.data_preparer = data_preparer
        self.chart_factory = ExcelChartFactory()

    def add_charts_sheet(
        self,
        writer: pd.ExcelWriter,
        summary_df: pd.DataFrame,
        stock_summary_df: pd.DataFrame,
    ) -> None:
        charts_ws = writer.sheets.get(SheetName.CHARTS)
        chart_ws = writer.sheets.get(SheetName.CHART_DATA)
        stock_ws = writer.sheets.get(SheetName.STOCK_SUMMARY)
        if charts_ws is None:
            return

        charts_ws._charts = []
        period_label = self.data_preparer.dashboard_stats(summary_df, stock_summary_df)["period"]
        self._style_charts_sheet(charts_ws, period_label)

        if chart_ws is not None and not summary_df.empty:
            self._add_summary_charts(charts_ws, chart_ws, summary_df)

        if stock_ws is not None and not stock_summary_df.empty:
            self._add_stock_charts(charts_ws, stock_ws, stock_summary_df)

        self._write_charts_footer(charts_ws)

    def _add_summary_charts(self, charts_ws, chart_ws, summary_df: pd.DataFrame) -> None:
        max_summary_row = len(summary_df) + 1
        categories = Reference(chart_ws, min_col=1, min_row=2, max_row=max_summary_row)
        tick_skip = max(1, len(summary_df) // 8)
        charts_ws.add_chart(
            self.chart_factory.cost_vs_value_chart(chart_ws, categories, max_summary_row, tick_skip),
            "B4",
        )
        charts_ws.add_chart(
            self.chart_factory.line_chart(
                title="Toplam Getiri Trendi (%)",
                y_axis_title="Getiri",
                data=Reference(chart_ws, min_col=5, max_col=5, min_row=1, max_row=max_summary_row),
                categories=categories,
                color=theme.GOLD_ACCENT,
                y_axis_format="0.00%",
                tick_skip=tick_skip,
                smooth=True,
            ),
            "B24",
        )
        charts_ws.add_chart(
            self.chart_factory.drawdown_chart(chart_ws, categories, max_summary_row, tick_skip),
            "N24",
        )

    def _add_stock_charts(self, charts_ws, stock_ws, stock_summary_df: pd.DataFrame) -> None:
        allocation_chart = self.chart_factory.allocation_doughnut(stock_ws, len(stock_summary_df) + 1)
        charts_ws.add_chart(allocation_chart, "N4")

    def _style_charts_sheet(self, worksheet, period_label: str | None = None) -> None:
        worksheet.sheet_view.showGridLines = False
        for column_idx in range(1, 26):
            worksheet.column_dimensions[get_column_letter(column_idx)].width = 12
        for row_idx in range(1, 90):
            worksheet.row_dimensions[row_idx].height = 20

        worksheet.merge_cells("A1:Z1")
        banner = worksheet["A1"]
        banner.value = "Portföy Performans Raporu — Grafikler"
        banner.font = Font(bold=True, color=theme.WHITE, size=16)
        banner.fill = PatternFill(start_color=theme.NAVY_PRIMARY, end_color=theme.NAVY_PRIMARY, fill_type="solid")
        banner.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 32

        worksheet.merge_cells("A2:Z2")
        sub = worksheet["A2"]
        sub.value = f"Rapor dönemi: {period_label}" if period_label and period_label != "Veri yok" else "Rapor dönemi: —"
        sub.font = Font(italic=True, color=theme.SLATE, size=10)
        sub.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[2].height = 22

        subtitles = {
            "B3":  "Portföy değeri ile alım maliyeti karşılaştırması",
            "B23": "Kümülatif getiri yüzdesi (%)",
            "N3":  "Hisse senedi ağırlık dağılımı",
            "N23": "Tepe değerden maksimum düşüş — risk ölçütü",
        }
        for cell_ref, text in subtitles.items():
            cell = worksheet[cell_ref]
            cell.value = text
            cell.font = Font(italic=True, color=theme.TEXT_MUTED, size=9)
            cell.alignment = Alignment(vertical="center")
        worksheet.row_dimensions[3].height = 18
        worksheet.row_dimensions[23].height = 18

    def _write_charts_footer(self, worksheet) -> None:
        footer_row = 85
        worksheet.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=26)
        cell = worksheet.cell(row=footer_row, column=1)
        cell.value = (
            "Veriler portföy işlem geçmişinden hesaplanmıştır. BIST kapalı günleri "
            "ve fiyat verisi olmayan günler performans serisine dahil edilmez."
        )
        cell.font = Font(italic=True, color=theme.TEXT_MUTED, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[footer_row].height = 28
